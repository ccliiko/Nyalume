"""工具注册表与纯函数工具。"""

from pathlib import Path
import threading

import pytest

from nyalume.core import memory, tools


def test_tool_schemas_include_reminder_tools():
    names = [t["function"]["name"] for t in tools.TOOL_SCHEMAS]
    assert "create_reminder" in names
    assert "remind_me_in" in names
    assert "cancel_reminder" in names
    assert "delete_note" in names
    assert "web_search" in names
    assert "pdf_edit" in names


def test_web_search_prefers_the_complete_title_and_drops_unrelated_baike(monkeypatch):
    calls = []

    def fake_search(term, _limit):
        calls.append(term)
        return [
            {
                "title": "千（汉语汉字）_百度百科",
                "url": "https://baike.baidu.com/item/%E5%8D%83/1896443",
                "snippet": "千，汉语一级字。",
            },
            {
                "title": "千恋万花 - 萌娘百科",
                "url": "https://example.com/senren-banka",
                "snippet": "《千恋万花》作品资料。",
            },
        ]

    monkeypatch.setattr(tools, "_bing_search", fake_search)
    result = tools._web_search("千恋万花")

    assert calls[0] == '"千恋万花"'
    assert "千恋万花 - 萌娘百科" in result
    assert "千（汉语汉字）" not in result


def test_remind_me_in_is_one_shot_not_every_minute():
    """“X 分钟后提醒一次”必须落成固定时刻 cron，不能变成 */1 每分钟。"""
    result = tools.execute_tool(
        "remind_me_in", {"content": "喝水", "minutes": 1}
    )
    assert "已设好一次性提醒" in result
    rows = tools.list_reminders()
    assert rows, "提醒应已入库"
    cron = rows[0]["cron"]
    assert "*/1" not in cron and cron.split()[0] != "*"
    assert cron.count(" ") == 4
    assert rows[0]["one_shot"] == 1, "remind_me_in 必须标记为一次性"


def test_create_reminder_stays_recurring():
    rows_before = tools.list_reminders()
    result = tools.execute_tool(
        "create_reminder", {"content": "周期性测试", "cron": "0 9 * * *"}
    )
    assert "已设置定时提醒" in result
    rows = tools.list_reminders()
    newest = rows[0]
    assert newest["one_shot"] == 0, "create_reminder 是周期提醒，不能标记一次性"
    tools.execute_tool("cancel_reminder", {"reminder_id": newest["id"]})


def test_calculator_safe_math():
    assert tools.execute_tool("calculator", {"expression": "(13*78)+5"}) == "1019"
    assert tools.execute_tool("calculator", {"expression": "2**10"}) == "1024"


def test_calculator_rejects_code():
    result = tools.execute_tool(
        "calculator", {"expression": "__import__('os').system('echo hi')"}
    )
    assert result.startswith("计算失败")


def test_note_save_and_list_with_tag():
    tag = "pytest-tag"
    memory.init_db()
    tools.execute_tool("save_note", {"content": "测试便签内容", "tag": tag})
    rows = memory.note_list(tag)
    assert "测试便签内容" in rows


def test_unknown_tool_returns_error_string():
    assert "未知工具" in tools.execute_tool("not_exist", {})


def test_doc_library_ingest_and_search(tmp_path):
    """把文件夹导入文档库后，可用关键词检索到原文片段。"""
    folder = tmp_path / "资料"
    folder.mkdir()
    (folder / "壁纸说明.txt").write_text(
        "桌面聊天窗壁纸支持不透明度设置", encoding="utf-8"
    )
    (folder / "顺手写的.md").write_text("这一篇与关键词无关", encoding="utf-8")

    result = tools.execute_tool("add_documents", {"path": str(folder)})
    assert "壁纸说明" in result and "已导入" in result
    assert "壁纸说明" in tools.execute_tool("list_docs", {})

    hits = memory.search_docs("不透明度")
    assert "壁纸说明" in hits and "不透明度" in hits
    assert "没有找到" in memory.search_docs("词库里没有的xyz")
    fuzzy = memory.search_docs("壁纸不透明度设多少好")
    assert "壁纸说明" in fuzzy
    hits = memory.retrieve_docs("桌面壁纸不透明度怎么设置", limit=2)
    assert hits and hits[0]["title"] == "壁纸说明.txt"


def test_doc_library_ingest_docx_and_xlsx(tmp_path):
    """Word/Excel 也能抽取文本入库（生成最小 docx/xlsx 再导入）。"""
    from docx import Document
    from openpyxl import Workbook

    doc_path = tmp_path / "说明.docx"
    doc = Document()
    doc.add_paragraph("Word 导入测试关键词")
    doc.save(doc_path)

    xl_path = tmp_path / "数据.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["项目", "金额"])
    ws.append(["Excel 导入测试关键词", 1014])
    wb.save(xl_path)

    result = tools.execute_tool("add_documents", {"path": str(tmp_path)})
    assert "说明.docx" in result and "数据.xlsx" in result
    assert "Word 导入测试关键词" in memory.search_docs("Word 导入测试关键词")
    assert "Excel 导入测试关键词" in memory.search_docs("Excel 导入测试关键词")


def test_scanned_pdf_without_text_layer_is_skipped(tmp_path):
    """只有空白页/重复水印的扫描 PDF 不能假装导入成功。"""
    from pypdf import PdfWriter

    pdf_path = tmp_path / "扫描件.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with open(pdf_path, "wb") as fh:
        writer.write(fh)

    result = tools.execute_tool("add_documents", {"path": str(pdf_path)})
    assert "导入失败" in result
    assert "扫描版" in result


def test_pdf_edit_page_operations(tmp_path, monkeypatch):
    """PDF 工具能合并、抽页、删页和旋转，且不覆盖已有文件。"""
    from pypdf import PdfReader, PdfWriter

    monkeypatch.setenv("AGENT_WORKSPACE", str(tmp_path))
    tools.set_permission_mode("workspace")
    first = PdfWriter()
    first.add_blank_page(width=100, height=200)
    first.add_blank_page(width=120, height=200)
    with open(tmp_path / "a.pdf", "wb") as fh:
        first.write(fh)
    second = PdfWriter()
    second.add_blank_page(width=300, height=200)
    with open(tmp_path / "b.pdf", "wb") as fh:
        second.write(fh)

    merged = tools.execute_tool(
        "pdf_edit",
        {
            "operation": "merge",
            "files": ["a.pdf", "b.pdf"],
            "output": "out/merged.pdf",
        },
    )
    assert "3 页" in merged
    assert len(PdfReader(tmp_path / "out" / "merged.pdf").pages) == 3

    extracted = tools.execute_tool(
        "pdf_edit",
        {
            "operation": "extract",
            "files": ["out/merged.pdf"],
            "output": "out/extracted.pdf",
            "pages": "3,1",
        },
    )
    assert "2 页" in extracted
    extracted_pages = PdfReader(tmp_path / "out" / "extracted.pdf").pages
    assert [float(page.mediabox.width) for page in extracted_pages] == [300.0, 100.0]

    deleted = tools.execute_tool(
        "pdf_edit",
        {
            "operation": "delete",
            "files": ["out/merged.pdf"],
            "output": "out/deleted.pdf",
            "pages": "2",
        },
    )
    assert "2 页" in deleted

    rotated = tools.execute_tool(
        "pdf_edit",
        {
            "operation": "rotate",
            "files": ["a.pdf"],
            "output": "out/rotated.pdf",
            "pages": "1",
            "degrees": 90,
        },
    )
    assert "2 页" in rotated
    rotated_pages = PdfReader(tmp_path / "out" / "rotated.pdf").pages
    assert rotated_pages[0].rotation == 90
    assert rotated_pages[1].rotation == 0

    exists = tools.execute_tool(
        "pdf_edit",
        {
            "operation": "merge",
            "files": ["a.pdf", "b.pdf"],
            "output": "out/merged.pdf",
        },
    )
    assert "输出已存在" in exists


def test_pdf_edit_obeys_permissions_and_exposes_paths(tmp_path, monkeypatch):
    monkeypatch.setenv("AGENT_WORKSPACE", str(tmp_path))
    tools.set_permission_mode("read_only")
    args = {
        "operation": "rotate",
        "files": ["a.pdf"],
        "output": "out.pdf",
    }
    assert "只读" in tools.execute_tool("pdf_edit", args)
    tools.set_permission_mode("workspace")
    assert tools.tool_paths("pdf_edit", args) == [
        str(tmp_path / "a.pdf"),
        str(tmp_path / "out.pdf"),
    ]
    outside = tmp_path.parent / "outside.pdf"
    need, reason = tools.approval_needed(
        "pdf_edit", {**args, "output": str(outside)}
    )
    assert need is True and "越出" in reason


def test_workspace_file_tools_roundtrip(tmp_path):
    """文件工具只在工作目录内生效：写/读/列表/移动/删除 + 越界拒绝。"""
    memory.set_setting("workspace_root", str(tmp_path / "项目"))
    assert "工作目录已设为" in tools.execute_tool(
        "set_workspace", {"path": str(tmp_path / "项目")}
    )
    assert memory.get_setting("workspace_root") == str(tmp_path / "项目")

    assert "已写入" in tools.execute_tool(
        "file_write", {"path": "src/main.py", "content": "print('hi')\n"}
    )
    assert "print('hi')" in tools.execute_tool("file_read", {"path": "src/main.py"})
    assert "main.py" in tools.execute_tool("file_list", {"path": "src"})
    assert "已移动" in tools.execute_tool(
        "file_move", {"src": "src/main.py", "dst": "main.py"}
    )
    assert "main.py" in tools.execute_tool("file_list", {})
    assert "已删除文件" in tools.execute_tool("file_delete", {"path": "main.py"})

    tools.execute_tool("file_write", {"path": "a/b.txt", "content": "x"})
    assert "删除失败" in tools.execute_tool("file_delete", {"path": "a"})  # 非空目录拒删
    assert "授权根" in tools.execute_tool("file_read", {"path": "..\\..\\agent.db"})


def test_plain_conversations_have_separate_named_workspaces(tmp_path, monkeypatch):
    monkeypatch.setenv("AGENT_WORKSPACE", str(tmp_path))
    first = memory.create_session()
    second = memory.create_session()
    memory.save_message(first, "user", "整理软件工程作业")
    memory.save_message(second, "user", "准备 Agent 面试")
    try:
        tools.set_session_context(first)
        tools.execute_tool("file_write", {"path": "答案.md", "content": "one"})
        first_root = tools.session_work_root(first)
        tools.clear_session_context()

        tools.set_session_context(second)
        tools.execute_tool("file_write", {"path": "答案.md", "content": "two"})
        second_root = tools.session_work_root(second)

        assert first_root != second_root
        assert "整理软件工程作业" in first_root
        assert "准备 Agent 面试" in second_root
        assert (Path(first_root) / "答案.md").read_text(encoding="utf-8") == "one"
        assert (Path(second_root) / "答案.md").read_text(encoding="utf-8") == "two"
    finally:
        tools.clear_session_context()
        memory.delete_session(first)
        memory.delete_session(second)


def test_parallel_conversations_do_not_share_tool_context(tmp_path, monkeypatch):
    monkeypatch.setenv("AGENT_WORKSPACE", str(tmp_path))
    first = memory.create_session()
    second = memory.create_session()
    memory.save_message(first, "user", "并发作业甲")
    memory.save_message(second, "user", "并发作业乙")
    barrier = threading.Barrier(2)
    results = {}

    def write_in_session(session_id: str, content: str) -> None:
        tools.reset_session_context(session_id)
        barrier.wait()
        results[session_id] = tools.execute_tool(
            "file_write", {"path": "同名.txt", "content": content}
        )
        tools.clear_session_context()

    workers = [
        threading.Thread(target=write_in_session, args=(first, "甲")),
        threading.Thread(target=write_in_session, args=(second, "乙")),
    ]
    try:
        for worker in workers:
            worker.start()
        for worker in workers:
            worker.join(timeout=5)
        assert all(not worker.is_alive() for worker in workers)
        assert all("已写入" in result for result in results.values())
        assert (Path(tools.session_work_root(first)) / "同名.txt").read_text(encoding="utf-8") == "甲"
        assert (Path(tools.session_work_root(second)) / "同名.txt").read_text(encoding="utf-8") == "乙"
    finally:
        tools.clear_session_context()
        memory.delete_session(first)
        memory.delete_session(second)


def test_session_temp_helpers_are_removed(tmp_path, monkeypatch):
    monkeypatch.setenv("AGENT_WORKSPACE", str(tmp_path))
    session_id = memory.create_session()
    memory.save_message(session_id, "user", "生成报告")
    try:
        tools.set_session_context(session_id)
        assert "已写入" in tools.execute_tool(
            "file_write", {"path": ".nyalume/tmp/helper.py", "content": "print(1)"}
        )
        root = Path(tools.session_work_root(session_id))
        assert (root / ".nyalume" / "tmp" / "helper.py").is_file()
        tools.cleanup_session_temp(session_id)
        assert not (root / ".nyalume" / "tmp").exists()
    finally:
        tools.clear_session_context()
        memory.delete_session(session_id)


def test_temp_cleanup_does_not_create_a_folder_for_chat_only_sessions(tmp_path, monkeypatch):
    monkeypatch.setenv("AGENT_WORKSPACE", str(tmp_path))
    session_id = memory.create_session()
    try:
        tools.cleanup_session_temp(session_id)
        assert not (tmp_path / "conversations").exists()
    finally:
        memory.delete_session(session_id)


def test_doc_library_ingest_pptx(tmp_path):
    """PPTX 幻灯片文字（标题/正文）能抽取入库。"""
    from pptx import Presentation

    pres = Presentation()
    slide = pres.slides.add_slide(pres.slide_layouts[1])
    slide.shapes.title.text = "PPT 导入测试标题"
    for placeholder in slide.placeholders:
        if placeholder != slide.shapes.title:
            placeholder.text = "信用卡逾期 Spark 数据处理"
            break
    file_path = tmp_path / "课程.pptx"
    pres.save(file_path)

    result = tools.execute_tool("add_documents", {"path": str(file_path)})
    assert "课程.pptx" in result
    assert "信用卡逾期 Spark" in memory.search_docs("信用卡逾期")


def test_project_multi_root_file_switch(tmp_path):
    """多根项目：file_set_root 只能切到授权目录，文件操作跟着根走。"""
    main = tmp_path / "主"
    extra = tmp_path / "附加"
    main.mkdir()
    extra.mkdir()
    project = memory.add_project(str(main))
    memory.add_project_folder(project["id"], str(extra))
    sid = memory.create_session(project=project["id"])
    tools.set_session_context(sid)
    try:
        assert "已切换操作根" in tools.execute_tool(
            "file_set_root", {"path": str(extra)}
        )
        assert "已写入" in tools.execute_tool(
            "file_write", {"path": "notes.md", "content": "附加根内容"}
        )
        assert "附加根内容" in tools.execute_tool("file_read", {"path": "notes.md"})
        outside = tmp_path / "外部"
        outside.mkdir()
        assert "不允许切到" in tools.execute_tool(
            "file_set_root", {"path": str(outside)}
        )
    finally:
        tools.clear_session_context()
        memory.delete_session(sid)
        memory.delete_project(project["id"])


def test_run_code_python_in_workspace(tmp_path):
    """受限执行器：在设定目录里跑 python，能看到 stdout 与退出码。"""
    memory.set_setting("workspace_root", str(tmp_path))
    assert "已写入" in tools.execute_tool(
        "file_write", {"path": "demo.py", "content": "print('执行器工作正常')"}
    )
    out = tools.execute_tool("run_code", {"command": "python demo.py"})
    assert "退出码 0" in out and "执行器工作正常" in out


def test_run_code_timeout_and_deny(tmp_path):
    """超时会终止；非白名单命令直接拒绝。"""
    memory.set_setting("workspace_root", str(tmp_path))
    slow = tools.execute_tool(
        "run_code",
        {"command": 'python -c "import time; time.sleep(5)"', "timeout": 1},
    )
    assert "超时" in slow
    denied = tools.execute_tool("run_code", {"command": "powershell whoami"})
    assert "只允许" in denied


def test_permission_modes_gate(tmp_path):
    """只读拦截写/运行；工作区锁根并给授权指引；全权允许用户点名路径。"""
    memory.set_setting("workspace_root", str(tmp_path))
    assert "已写入" in tools.execute_tool(
        "file_write", {"path": "a.txt", "content": "x"}
    )

    tools.set_permission_mode("read_only")
    assert "只读" in tools.execute_tool(
        "file_write", {"path": "b.txt", "content": "x"}
    )
    assert "只读" in tools.execute_tool(
        "run_code", {"command": 'python -c "print(1)"'}
    )
    assert "x" in tools.execute_tool("file_read", {"path": "a.txt"})

    tools.set_permission_mode("workspace")

    tools.set_permission_mode("daily")
    assert "不会调用工具" in tools.execute_tool("get_current_time", {})
    tools.set_permission_mode("workspace")
    outside = tmp_path.parent / ("outside_" + tmp_path.name + ".txt")
    out = tools.execute_tool("file_write", {"path": str(outside), "content": "no"})
    assert "全权" in out and "授权" in out  # 越界写被拦并提示升级选项

    tools.set_permission_mode("full")
    out2 = tools.execute_tool("file_write", {"path": str(outside), "content": "full ok"})
    assert "已写入" in out2
    assert "系统目录" in tools.execute_tool("file_delete", {"path": r"C:\Windows\win.ini"})
    tools.execute_tool("file_delete", {"path": str(outside)})
    tools.set_permission_mode("workspace")


def test_approval_needed_and_undo(tmp_path):
    """审批判定：根内文件不用问；越界/只读要问；撤销能还原文本。"""
    memory.set_setting("workspace_root", str(tmp_path))
    tools.set_permission_mode("workspace")
    need, _ = tools.approval_needed("file_write", {"path": "a.txt"})
    assert need is False
    outside = tmp_path.parent / ("need_" + tmp_path.name + ".txt")
    need2, desc = tools.approval_needed("file_write", {"path": str(outside)})
    assert need2 is True and "越出" in desc
    tools.set_permission_mode("full")
    assert tools.approval_needed("file_write", {"path": str(outside)}) == (False, "")
    tools.set_permission_mode("workspace")

    f = tmp_path / "u.txt"
    f.write_text("v2", encoding="utf-8")
    memory.add_undo("undo-real", "s1", "write", str(f), "v1")
    assert "已撤销" in tools.apply_undo("undo-real")
    assert f.read_text(encoding="utf-8") == "v1"


def test_approval_remember_rule(tmp_path):
    """同意并记住后，同会话同类越界请求不再询问。"""
    memory.set_setting("workspace_root", str(tmp_path))
    outside = tmp_path.parent / ("remember_" + tmp_path.name + ".txt")
    tools.set_session_context("sess-remember")
    try:
        tools.set_permission_mode("workspace")
        need, _ = tools.approval_needed("file_write", {"path": str(outside)})
        assert need is True
        tools.record_approval_rule("file_write")
        assert tools.auto_approved("file_write") is True
        need2, _ = tools.approval_needed("file_write", {"path": str(outside)})
        assert need2 is False
    finally:
        tools.clear_session_context()
        tools.set_permission_mode("workspace")


def test_diff_undo_counts(tmp_path):
    """diff_undo 输出 +/− 行数与按块（含一行上下文）的对比。"""
    target = tmp_path / "d.py"
    target.write_text("a\nb\nc\n", encoding="utf-8")
    memory.add_undo("undo-diff-1", "s2", "write", str(target), "a\nb\nc\n")
    target.write_text("a\nx\nc\nd\n", encoding="utf-8")
    diff = tools.diff_undo("undo-diff-1")
    assert diff["added"] == 2 and diff["removed"] == 1
    tags = [tag for block in diff["hunks"] for tag, _ in block]
    assert "add" in tags and "del" in tags and "ctx" in tags


def test_project_inventory_lists_files(tmp_path):
    """项目清单列出文件、跳过噪音目录，供首次对话主动浏览。"""
    root = tmp_path / "项目"
    (root / "src").mkdir(parents=True)
    (root / "README.md").write_text("# 项目", encoding="utf-8")
    (root / "src" / "main.py").write_text("print(1)", encoding="utf-8")
    (root / ".venv" / "x.py").mkdir(parents=True)
    (root / ".venv" / "x.py" / "junk.txt").write_text("", encoding="utf-8")
    project = memory.add_project(str(root))
    inventory = tools.project_inventory(project["id"])
    assert "README.md" in inventory and "main.py" in inventory
    assert "junk.txt" not in inventory
    memory.delete_project(project["id"])


def test_tool_paths_are_absolute_and_follow_session_project(tmp_path):
    """前端拿到绝对路径，点击时才能准确打开项目内的文件。"""
    project_root = tmp_path / "课程项目"
    project_root.mkdir()
    project = memory.add_project(str(project_root))
    session_id = memory.create_session(project=project["id"])
    tools.set_session_context(session_id)
    try:
        assert tools.tool_paths("file_write", {"path": "报告.pdf"}) == [
            str(project_root / "报告.pdf")
        ]
        assert tools.tool_paths("run_code", {"cwd": "scripts"}) == [
            str(project_root / "scripts")
        ]
        assert tools.tool_paths("add_documents", {"path": "资料"}) == [
            str(project_root / "资料")
        ]
    finally:
        tools.clear_session_context()
        memory.delete_session(session_id)
        memory.delete_project(project["id"])


def test_search_files_scope_and_approval(tmp_path, monkeypatch):
    """search_files：授权根内直接搜；根外要审批；全权模式放行。"""
    import os as _os

    monkeypatch.setenv("AGENT_WORKSPACE", str(tmp_path))
    memory.set_setting("workspace_root", str(tmp_path))
    (tmp_path / "小夏.txt").write_text("x", encoding="utf-8")
    sub = tmp_path / "子"
    sub.mkdir()
    (sub / "陆景和.md").write_text("x", encoding="utf-8")
    (sub / "青鸟文件夹").mkdir()
    tools.set_permission_mode("workspace")

    inside = tools.execute_tool("search_files", {"keyword": "小夏"})
    assert "小夏.txt" in inside
    glob = tools.execute_tool("search_files", {"keyword": "*.md"})
    assert "陆景和.md" in glob
    dirhit = tools.execute_tool("search_files", {"keyword": "青鸟"})
    assert "[目录] " in dirhit and "青鸟文件夹" in dirhit

    outside_dir = tmp_path.parent / ("search_out_" + tmp_path.name)
    outside_dir.mkdir(exist_ok=True)
    outside = outside_dir / "outside_target.txt"
    outside.write_text("x", encoding="utf-8")
    try:
        denied = tools.execute_tool(
            "search_files", {"keyword": "outside_target", "base": str(outside_dir)}
        )
        assert "授权根" in denied
        need, _ = tools.approval_needed(
            "search_files", {"keyword": "outside_target", "base": str(outside_dir)}
        )
        assert need is True
        tools.set_permission_mode("full")
        allowed = tools.execute_tool(
            "search_files", {"keyword": "outside_target", "base": str(outside_dir)}
        )
        assert "outside_target.txt" in allowed
    finally:
        _os.remove(outside)
        outside_dir.rmdir()
        tools.set_permission_mode("workspace")
    memory.set_setting("workspace_root", "")


def test_loop_breaker_blocks_identical_repeats():
    """纪律闸门：同一轮相同(工具, 参数)重复 3 次被拦，换参数或重置后放行。"""
    tools.reset_hook_state()
    arg = {"expression": "7*8"}
    assert "56" in tools.execute_tool("calculator", arg)
    assert "56" in tools.execute_tool("calculator", arg)
    blocked = tools.execute_tool("calculator", arg)
    assert "纪律闸门" in blocked and "重复调用" in blocked
    # 不同参数不算重复
    assert "9" in tools.execute_tool("calculator", {"expression": "3*3"})
    # 新的一轮重置后同参数可再用
    tools.reset_hook_state()
    assert "56" in tools.execute_tool("calculator", arg)
